# -*- coding: utf-8 -*-
# 台帳＋AgiMaレポートから prices.json を生成（型式→素体/PSA10）
#
# 使い方:
#   python gen_prices.py            … 1回生成
#   python gen_prices.py --push     … 生成して変化があれば git commit & push（アプリに即反映）
#   python gen_prices.py --watch    … 台帳/レポートの保存を監視し、更新されるたび生成＋push
#                                     （--interval 秒 で監視間隔、既定60秒）
# パスは環境変数 AGIS_DATA / AGIS_OUT でも上書きできる。
import openpyxl, glob, os, re, json, datetime, argparse, subprocess, time, sys

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.environ.get("AGIS_DATA", r"C:\Users\PC_User\Desktop\ClaudeCode\data\Agito-Market")
OUT = os.environ.get("AGIS_OUT", os.path.join(HERE, "prices.json"))

def candidates(pat):
    return [f for f in glob.glob(os.path.join(DATA, pat)) if not os.path.basename(f).startswith("~$")]

def latest(pat):
    # ファイル名の日付(YYYYMMDD)優先、同日/日付なしは更新時刻で決める
    fs = candidates(pat)
    def key(f):
        m = re.search(r"(\d{8})", os.path.basename(f))
        return (m.group(1) if m else "0", os.path.getmtime(f))
    return max(fs, key=key) if fs else None

def slug(t):
    s = re.sub(r"[\s/]+", "-", str(t).strip().lower())
    s = re.sub(r"-+", "-", s).strip("-")
    return s

def num(v):
    return v if isinstance(v, (int, float)) else None

def build():
    db = {}

    def put(k, name, sotai, psa10, src, overwrite):
        if not k: return
        cur = db.get(k)
        if cur and not overwrite and cur.get("psa10"): return
        if cur and overwrite:
            # 台帳は下の行ほど新しい。新しい行の値を優先し、空欄は古い行の値で埋める
            sotai = sotai if sotai is not None else cur.get("sotai")
            psa10 = psa10 if psa10 is not None else cur.get("psa10")
        db[k] = {"name": name, "sotai": sotai, "psa10": psa10, "src": src}

    # 台帳（優先）
    lg = latest("*ledger*.xlsx")
    if lg:
        wb = openpyxl.load_workbook(lg, data_only=True, read_only=True)
        sg = [w for w in wb.worksheets if "シングル" in w.title][0]
        for row in sg.iter_rows(min_row=2, values_only=True):
            row = list(row) + [None] * 24
            nm, typ = row[2], row[3]
            if not nm: continue
            nm = str(nm).strip()
            # 型式未登録の直近カードも名前で引けるように収録（キーは noid-名前）
            k = slug(typ) if typ else "noid-" + slug(nm)
            put(k, nm, num(row[22]), num(row[23]), "台帳", True)
        wb.close()

    # AgiMaレポート ③乖離ピック / ②TOP / ①差分
    rp = latest("AgiMa相場レポート_*.xlsx")
    if rp:
        wb = openpyxl.load_workbook(rp, data_only=True, read_only=True)
        for ws in wb.worksheets:
            t = ws.title
            for row in ws.iter_rows(min_row=2, values_only=True):
                row = list(row) + [None] * 5
                a = row[0]
                if a is None or str(a).startswith("【"): continue
                if not row[2]: continue
                k = slug(row[2])
                if "乖離" in t:   # 判定|名|型番|美品|PSA10
                    put(k, row[1], num(row[3]), num(row[4]), "GR乖離", False)
                elif "TOP" in t or "取引" in t:  # 順位|銘柄|型番|件数|PSA10
                    put(k, row[1], None, num(row[4]), "GR出来高", False)
                elif "差分" in t:  # ID|名|型式|旧|新PSA10
                    put(k, row[1], None, num(row[4]), "GR差分", False)
        wb.close()

    # 型式が後から登録されたカードは noid 側を消す（同名の重複表示を防ぐ）
    named = {v["name"] for k, v in db.items() if not k.startswith("noid-")}
    for k in [k for k, v in db.items() if k.startswith("noid-") and v["name"] in named]:
        del db[k]

    return db, [f for f in (lg, rp) if f]

def write(db, srcs):
    now = datetime.datetime.now().replace(microsecond=0)
    out = {"updated": now.isoformat(), "count": len(db),
           "sources": [os.path.basename(f) for f in srcs], "cards": db}
    old = None
    if os.path.exists(OUT):
        try:
            with open(OUT, encoding="utf-8") as f: old = json.load(f)
        except Exception: pass
    if old and old.get("cards") == db and old.get("sources") == out["sources"]:
        print("prices.json: 変化なし（%d件）" % len(db))
        return False
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=0)
    print("prices.json:", len(db), "件 ->", OUT, "(", out["updated"], ")")
    return True

def push():
    repo = os.path.dirname(os.path.abspath(OUT))
    def git(*a):
        return subprocess.run(["git", "-C", repo] + list(a), capture_output=True, text=True)
    git("add", os.path.basename(OUT))
    if git("diff", "--cached", "--quiet").returncode == 0:
        return
    msg = "prices: 相場データ更新 " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    r = git("commit", "-m", msg)
    if r.returncode != 0:
        print("git commit 失敗:", r.stderr.strip()); return
    for wait in (0, 2, 4, 8, 16):
        time.sleep(wait)
        r = git("push")
        if r.returncode == 0:
            print("push 完了"); return
    print("git push 失敗:", r.stderr.strip())

def run(do_push):
    db, srcs = build()
    if write(db, srcs) and do_push:
        push()

def stamp():
    fs = candidates("*ledger*.xlsx") + candidates("AgiMa相場レポート_*.xlsx")
    return tuple(sorted((f, os.path.getmtime(f)) for f in fs))

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--push", action="store_true", help="変化があれば git commit & push")
    ap.add_argument("--watch", action="store_true", help="台帳/レポートの更新を監視して自動生成（push込み）")
    ap.add_argument("--interval", type=int, default=60, help="監視間隔（秒）")
    a = ap.parse_args()
    if not a.watch:
        run(a.push); sys.exit(0)
    print("監視中:", DATA, "（Ctrl+Cで終了）")
    last = None
    while True:
        try:
            s = stamp()
            if s != last:
                time.sleep(3)  # Excel保存直後の書き込み完了を待つ
                run(True)
                last = stamp()
        except Exception as e:  # 保存中のロック等は次の周回で再試行
            print("スキップ:", e)
        time.sleep(a.interval)
